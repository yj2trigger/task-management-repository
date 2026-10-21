"""원본 강의자료(PDF·PPTX·PPT·DOCX)에서 텍스트를 뽑는다. 원문 대조(verify.py)가 쓴다.

원본은 사용자가 대화에 올린 파일이라 UPLOADS 폴더에 있다(파일 이름 앞에 해시가 붙어 있어 접두사로 찾는다).
"""
import glob, os, struct

UPLOADS = os.path.join(os.environ["USERPROFILE"], ".claude", "uploads", "f82f0f52-550e-4d75-a7a3-0b05f6c42a0a")
FILES = {  # 인용 출처 이름 → 업로드 파일 접두사
    "1주 PPT": "1411303b", "2강 PPT": "d2496d2a", "2강 워크시트": "855b808e", "3강 PPT": "df1ab502",
    "3강 워크시트": "f1c55038", "5강 PPT": "5d1d21c6", "5강 워크시트": "d8f196ab", "6강 교재": "46fcfdff",
    "6강 워크시트": "5e6aded2", "정리자료": "1b137db6", "답안 작성법": "2f80e9d5",
}


def _ppt(path):
    import olefile
    data = olefile.OleFileIO(path).openstream("PowerPoint Document").read()
    out = []

    def walk(start, end):
        p = start
        while p + 8 <= end:
            ver, typ, ln = struct.unpack("<HHI", data[p:p + 8])
            body = p + 8
            if ver & 0xF == 0xF:
                walk(body, min(body + ln, end))
            elif typ == 4000:
                out.append(data[body:body + ln].decode("utf-16le", "ignore"))
            elif typ == 4008:
                out.append(data[body:body + ln].decode("cp949", "ignore"))
            p = body + ln
    walk(0, len(data))
    return "\n".join(out)


def _docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []
    for ch in d.element.body.iterchildren():
        if ch.tag.endswith("}p"):
            out.append(Paragraph(ch, d).text)
        elif ch.tag.endswith("}tbl"):
            out += [" | ".join(c.text for c in r.cells) for r in Table(ch, d).rows]
    return "\n".join(out)


def _pptx(path):
    from pptx import Presentation
    out = []
    for sl in Presentation(path).slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                out += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
    return "\n".join(out)


def _pdf(path):
    import pymupdf
    return "".join(p.get_text() for p in pymupdf.open(path))


def load(name):
    path = glob.glob(os.path.join(UPLOADS, FILES[name] + "*"))[0]
    ext = path.rsplit(".", 1)[-1].lower()
    return {"pdf": _pdf, "docx": _docx, "pptx": _pptx, "ppt": _ppt}[ext](path)
